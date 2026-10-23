"""O1 隐藏验收：Q1 汇总 + 明细，数值、格式、合计都要对。"""

import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "out" / "quarterly.xlsx"


def sheet_rows(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


class QuarterlyAccept(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        assert XLSX.exists(), f"没有找到产物: {XLSX}"
        cls.wb = load_workbook(XLSX)

    def test_two_sheets_named_exactly(self):
        self.assertEqual(self.wb.sheetnames, ["Q1", "明细"])

    def test_q1_header(self):
        rows = sheet_rows(self.wb["Q1"])
        self.assertEqual(rows[0][:4], ["月份", "订单数", "销售额(元)", "环比"])

    def test_month_rows_and_counts(self):
        rows = sheet_rows(self.wb["Q1"])
        body = {r[0]: r for r in rows[1:] if r[0] in ("2026-01", "2026-02", "2026-03")}
        self.assertEqual(list(body), ["2026-01", "2026-02", "2026-03"])
        for month in body:
            self.assertEqual(body[month][1], 2, f"{month} 订单数不对")

    def test_month_revenue(self):
        rows = sheet_rows(self.wb["Q1"])
        body = {r[0]: r for r in rows[1:]}
        expected = {"2026-01": 1900.0, "2026-02": 1220.0, "2026-03": 1722.0}
        for month, want in expected.items():
            got = body[month][2]
            self.assertIsInstance(got, (int, float), f"{month} 销售额不是数值: {got!r}")
            self.assertAlmostEqual(float(got), want, places=2, msg=f"{month} 销售额")

    def test_mom_values_and_percent_format(self):
        ws = self.wb["Q1"]
        rows = sheet_rows(ws)
        header = rows[0]
        col = header.index("环比")
        body = {r[0]: i for i, r in enumerate(rows)}
        # 第一行是 "—"
        self.assertEqual(ws.cell(row=body["2026-01"] + 1, column=col + 1).value, "—")
        for month, want in (("2026-02", -0.358), ("2026-03", 0.411)):
            cell = ws.cell(row=body[month] + 1, column=col + 1)
            self.assertIsInstance(cell.value, (int, float), f"{month} 环比不是数值")
            self.assertAlmostEqual(float(cell.value), want, delta=0.002, msg=f"{month} 环比")
            self.assertIn("%", cell.number_format or "", f"{month} 环比没有百分比格式")

    def test_total_row(self):
        rows = sheet_rows(self.wb["Q1"])
        total = next((r for r in rows[1:] if r[0] == "合计"), None)
        self.assertIsNotNone(total, "Q1 缺少合计行")
        self.assertEqual(total[1], 6)
        self.assertAlmostEqual(float(total[2]), 4842.0, places=2)

    def test_detail_header_and_rows(self):
        rows = sheet_rows(self.wb["明细"])
        self.assertEqual(rows[0][:4], ["sku", "month", "units", "销售额(元)"])
        self.assertEqual(len(rows) - 1, 6)
        self.assertEqual([r[0] for r in rows[1:]], ["A-1", "A-2", "B-1", "B-2", "C-1", "C-2"])

    def test_detail_revenue_matches_total(self):
        rows = sheet_rows(self.wb["明细"])
        values = []
        for r in rows[1:]:
            self.assertIsInstance(r[3], (int, float), f"{r[0]} 销售额不是数值")
            values.append(float(r[3]))
        self.assertAlmostEqual(sum(values), 4842.0, places=2)
        self.assertAlmostEqual(values[0], 900.0, places=2)


if __name__ == "__main__":
    unittest.main()
