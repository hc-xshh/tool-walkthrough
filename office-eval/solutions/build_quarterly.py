"""参考产物：按 spec.md 生成 out/quarterly.xlsx。

带 --variant 可以故意做错，用来验证判分口径：
  correct / missing_detail / wrong_total / string_numbers

用法：python solutions/build_quarterly.py --workdir <repo> [--variant correct]
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook

HEAD_Q1 = ["月份", "订单数", "销售额(元)", "环比"]
HEAD_DETAIL = ["sku", "month", "units", "销售额(元)"]


def row_revenue(row):
    return int(row["units"]) * float(row["unit_price"]) * (1 - float(row["discount"]))


def build(workdir, variant="correct"):
    workdir = Path(workdir)
    rows = list(csv.DictReader((workdir / "data" / "sales.csv").open(encoding="utf-8")))
    months = []
    detail = []
    for row in rows:
        rev = row_revenue(row)
        detail.append((row["sku"], row["month"], int(row["units"]), rev))
        if not months or months[-1][0] != row["month"]:
            months.append([row["month"], 0, 0.0])

    by_month = {}
    for month, units, _ in months:
        picked = [d for d in detail if d[1] == month]
        by_month[month] = [len(picked), sum(d[3] for d in picked)]

    wb = Workbook()
    q1 = wb.active
    q1.title = "Q1"
    q1.append(HEAD_Q1)

    previous = None
    for month, _, _ in months:
        count, revenue = by_month[month]
        if previous is None:
            mom = "—"
        else:
            mom = (revenue - previous) / previous
        cell = q1.cell(row=q1.max_row + 1, column=1, value=month)
        q1.cell(row=cell.row, column=2, value=count)
        q1.cell(row=cell.row, column=3, value=round(revenue, 2))
        mom_cell = q1.cell(row=cell.row, column=4, value=mom)
        if mom != "—":
            mom_cell.number_format = "0.0%"
        previous = revenue

    total_revenue = sum(v[1] for v in by_month.values())
    if variant == "wrong_total":
        total_revenue += 100
    q1.append(["合计", len(detail), round(total_revenue, 2), None])

    if variant != "missing_detail":
        ws = wb.create_sheet("明细")
        ws.append(HEAD_DETAIL)
        for sku, month, units, revenue in detail:
            value = round(revenue, 2)
            if variant == "string_numbers":
                value = f"{value:.2f}"
            ws.append([sku, month, units, value])

    out = workdir / "out" / "quarterly.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir", required=True)
    ap.add_argument("--variant", default="correct",
                    choices=["correct", "missing_detail", "wrong_total", "string_numbers"])
    a = ap.parse_args()
    print(build(a.workdir, a.variant))


if __name__ == "__main__":
    main()
