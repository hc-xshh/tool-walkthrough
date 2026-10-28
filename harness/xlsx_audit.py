#!/usr/bin/env python3
"""Audit an evidence workbook for the pitfalls hit in the 2026-10 walkthrough.

Rules:
- no image may be reused in more than one cell (hash-based)
- image must fit its anchor cell (width/height in px vs column/row)
- Excel row height must not exceed 409.5 pt (hard limit)
- status 未查到/无 must carry a 【缺失证据】 caption
- status 有 must not carry a 【缺失证据】 caption
- optional: expected number of images

Usage:
    python harness/xlsx_audit.py book.xlsx --sheet "Codex 45格" --expect-images 45
Exit code 1 on issues.
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import EMU_to_pixels

MAX_ROW_HEIGHT_PT = 409.5
DEFAULT_COL_WIDTH = 8.43
DEFAULT_ROW_HEIGHT_PT = 15.0


@dataclass
class Issue:
    code: str
    cell: str
    message: str


def _col_px(width_chars: float) -> int:
    # Excel's default font approximation: 7 px per character + 5 px padding.
    return int(round(width_chars * 7)) + 5


def _hash(path: Path) -> str | None:
    try:
        return hashlib.sha256(path.read_bytes()).hexdigest()
    except OSError:
        return None


def audit_workbook(path: Path, sheet: str | None = None, expect_images: int | None = None) -> list[Issue]:
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    issues: list[Issue] = []
    seen: dict[str, str] = {}

    for img in ws._images:
        anchor = img.anchor
        if not hasattr(anchor, "_from"):
            issues.append(Issue("unanchored", "?", "image has no one-cell anchor"))
            continue
        row = anchor._from.row + 1
        col = anchor._from.col + 1
        cell = f"{get_column_letter(col)}{row}"
        width_px = int(img.width or EMU_to_pixels(anchor.ext.cx))
        height_px = int(img.height or EMU_to_pixels(anchor.ext.cy))

        row_h = ws.row_dimensions[row].height or DEFAULT_ROW_HEIGHT_PT
        col_w = ws.column_dimensions[get_column_letter(col)].width or DEFAULT_COL_WIDTH
        max_h = row_h * 96 / 72
        max_w = _col_px(col_w)

        if row_h > MAX_ROW_HEIGHT_PT:
            issues.append(Issue("row-too-tall", cell, f"row height {row_h}pt > {MAX_ROW_HEIGHT_PT}pt"))
        if height_px > max_h + 1:
            issues.append(Issue("image-too-tall", cell, f"image {height_px}px > row {max_h:.0f}px ({row_h}pt)"))
        if width_px > max_w + 1:
            issues.append(Issue("image-too-wide", cell, f"image {width_px}px > col {max_w}px ({col_w} chars)"))

        ref = getattr(img, "ref", None)
        data = None
        if hasattr(ref, "getvalue"):
            data = ref.getvalue()
        elif getattr(img, "path", None):
            try:
                data = Path(img.path).read_bytes()
            except OSError:
                data = None
        if data:
            h = hashlib.sha256(data).hexdigest()
            if h in seen and seen[h] != cell:
                issues.append(Issue("duplicate-image", cell, f"same image also used at {seen[h]}"))
            seen.setdefault(h, cell)

        status = str(ws.cell(row, max(1, col - 1)).value or "").split("｜", 1)[0].strip()
        caption = str(ws.cell(row, col + 1).value or "")
        if status in ("未查到", "无") and "缺失证据" not in caption:
            issues.append(Issue("missing-absence-label", cell, f"status {status} caption must start with 【缺失证据】"))
        if status == "有" and "缺失证据" in caption:
            issues.append(Issue("false-absence-label", cell, "status 有 cannot use 【缺失证据】 caption"))

    if expect_images is not None and len(ws._images) != expect_images:
        issues.append(Issue("image-count", "sheet", f"{len(ws._images)} images, expected {expect_images}"))
    return issues


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--sheet")
    ap.add_argument("--expect-images", type=int)
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args(argv)

    issues = audit_workbook(args.workbook, args.sheet, args.expect_images)
    if args.json:
        import json
        print(json.dumps([i.__dict__ for i in issues], ensure_ascii=False, indent=2))
    else:
        for i in issues:
            print(f"{i.code:22} {i.cell:8} {i.message}")
        print(f"{len(issues)} issue(s)")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
