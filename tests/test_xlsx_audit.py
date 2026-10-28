"""xlsx_audit: duplicate images, cell overflow, absence captions."""
import pathlib
import sys
import tempfile
import unittest

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "harness"))

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

import xlsx_audit

try:
    from PIL import Image
    HAVE_PIL = True
except ImportError:
    HAVE_PIL = False


@unittest.skipUnless(HAVE_PIL, "Pillow not installed")
class XlsxAuditTest(unittest.TestCase):
    def _book(self, root, rows):
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 132
        ws.column_dimensions["D"].width = 60
        for i, (name, status, caption, img_path, h) in enumerate(rows, 1):
            r = i + 3
            ws.row_dimensions[r].height = h
            ws.cell(r, 1, name)
            ws.cell(r, 2, f"{status} ｜ note")
            ws.cell(r, 4, caption)
            img = XLImage(str(img_path))
            img.width, img.height = 900, 360
            anchor = OneCellAnchor()
            anchor._from = AnchorMarker(col=2, row=r - 1, colOff=pixels_to_EMU(3), rowOff=pixels_to_EMU(3))
            anchor.ext = XDRPositiveSize2D(pixels_to_EMU(900), pixels_to_EMU(360))
            img.anchor = anchor
            ws.add_image(img)
        p = root / "book.xlsx"
        wb.save(p)
        return p

    def test_clean_book(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("item", "有", "ok", p, 300)])
            self.assertEqual(xlsx_audit.audit_workbook(book, "S", 1), [])

    def test_duplicate_image(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("a", "有", "ok", p, 300), ("b", "有", "ok", p, 300)])
            codes = [i.code for i in xlsx_audit.audit_workbook(book, "S")]
            self.assertIn("duplicate-image", codes)

    def test_row_height_limit(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("a", "有", "ok", p, 420)])
            codes = [i.code for i in xlsx_audit.audit_workbook(book, "S")]
            self.assertIn("row-too-tall", codes)

    def test_image_too_tall(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("a", "有", "ok", p, 250)])
            codes = [i.code for i in xlsx_audit.audit_workbook(book, "S")]
            self.assertIn("image-too-tall", codes)

    def test_absence_caption(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("a", "未查到", "help 无此页", p, 300)])
            codes = [i.code for i in xlsx_audit.audit_workbook(book, "S")]
            self.assertIn("missing-absence-label", codes)

    def test_false_absence_label(self):
        with tempfile.TemporaryDirectory() as d:
            root = pathlib.Path(d)
            p = root / "a.png"
            Image.new("RGB", (900, 360), "white").save(p)
            book = self._book(root, [("a", "有", "【缺失证据】ok", p, 300)])
            codes = [i.code for i in xlsx_audit.audit_workbook(book, "S")]
            self.assertIn("false-absence-label", codes)


if __name__ == "__main__":
    unittest.main()
